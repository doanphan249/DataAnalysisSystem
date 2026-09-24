"""Chạy: python -m unittest discover -s src/tests -v (không cần cài pytest)."""
import contextlib
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'src'))
import openpyxl
import xlsxwriter
from reporting.common import classify_product, country_name, normalize_product, rounded
from reporting.models import DataError, Report, Settings
from reporting.sources import discover, Source
from reporting.pipeline import load_config, run_report
from reporting.excel_writer import write_report
from reporting.validation import validate
from reporting import receivables, sales_summary


class DiscoveryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def month(self, month, year=2026):
        folder = self.root / f'BCQT_{month:02d}{year}'
        folder.mkdir()
        for name in ['Report QTSX.xlsx', 'Report KQ Tieu thu.xlsx', '~$Report QTSX.xlsx']:
            (folder / name).touch()
        return folder

    def test_missing_month_rejected(self):
        self.month(1)
        self.month(3)
        with self.assertRaisesRegex(DataError, 'đủ tháng'):
            discover(self.root, None)

    def test_duplicate_source_rejected(self):
        folder = self.month(1)
        (folder / 'Another QTSX.xlsx').touch()
        with self.assertRaisesRegex(DataError, 'đúng 1 file'):
            discover(self.root, None)

    def test_lock_file_ignored(self):
        self.month(1)
        self.assertEqual(len(discover(self.root, None)), 1)

    def test_year_must_be_selected_when_ambiguous(self):
        self.month(1, 2026)
        self.month(1, 2027)
        with self.assertRaisesRegex(DataError, 'REPORT_YEAR'):
            discover(self.root, None)
        self.assertEqual(discover(self.root, 2027)[0].year, 2027)


class DataTests(unittest.TestCase):
    def test_negative_rounding_and_classification(self):
        self.assertEqual(rounded(-1.5), -1)
        self.assertEqual(classify_product('26/1 CVCCmP (60/40R) W'),
                         ('Recycle', 'Sợi CVC', 'Sợi chải kỹ', 26))
        self.assertEqual(classify_product('30/1 CVCCd (60/40) W')[2], 'Sợi chải thô')

    def test_unknown_name_does_not_invent_blend(self):
        self.assertEqual(normalize_product('99/1 Unknown (50/50re)-w', {}),
                         '99/1 Unknown (50/50re) W')

    def test_duplicate_key_is_error(self):
        report = Report(load_config('schemas.json'))
        report.rows['KyQuanTri'] = [['T1', 'Tháng 1'], ['T1', 'Tháng 1']]
        with self.assertRaisesRegex(DataError, 'Trùng khóa'):
            validate(report)

    def test_source_error_not_silently_zero(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'source.xlsx'
            with xlsxwriter.Workbook(path) as workbook:
                sheet = workbook.add_worksheet('Test')
                sheet.write_formula('A1', '=1/0', None, '#DIV/0!')
                sheet.write_string('A2', '\n')
            with Source(path, 2026, 1, Report({})) as source:
                with self.assertRaises(DataError):
                    source.number('Test', 1, 1, blank_zero=True)
                self.assertEqual(source.number('Test', 2, 1, blank_zero=True), 0)

    def test_formula_without_cache_is_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'source.xlsx'
            workbook = openpyxl.Workbook()
            workbook.active.title = 'Test'
            workbook.active['A1'] = '=1+2'
            workbook.save(path)
            workbook.close()
            with Source(path, 2026, 1, Report({})) as source:
                with self.assertRaisesRegex(DataError, 'chưa có giá trị'):
                    source.number('Test', 1, 1, blank_zero=True)

    def test_atomic_failure_preserves_old_file(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'output.xlsx'
            path.write_bytes(b'original-file')
            report = Report({'KyQuanTri': ['Ky', 'KyQuanTri']})
            report.add('KyQuanTri', ['T1', 'Tháng 1'], 'test')
            with patch('reporting.excel_writer.os.replace', side_effect=PermissionError('locked')):
                with self.assertRaises(DataError):
                    write_report(report, path, 'test')
            self.assertEqual(path.read_bytes(), b'original-file')
            self.assertEqual((path.parent / 'backups/output_test.xlsx').read_bytes(), b'original-file')
            self.assertEqual(list(path.parent.glob('.report_*.xlsx')), [])

    def test_empty_sheet_still_has_table(self):
        # Power BI đọc theo tên bảng: năm không có hợp đồng vẫn phải có MuaBongTable.
        schemas = load_config('schemas.json')
        with tempfile.TemporaryDirectory() as temp:
            report = Report({'KyQuanTri': schemas['KyQuanTri'], 'HDBong': schemas['HDBong']})
            report.add('KyQuanTri', ['T1', 'Tháng 1'], 'test')
            path = Path(temp) / 'result.xlsx'
            write_report(report, path, 'test')
            workbook = openpyxl.load_workbook(path)
            self.assertIn('MuaBongTable', workbook['HDBong'].tables)
            workbook.close()

    def test_formula_like_text_is_text(self):
        with tempfile.TemporaryDirectory() as temp:
            report = Report({'KyQuanTri': ['Ky', 'KyQuanTri']})
            report.add('KyQuanTri', ['T1', '=HYPERLINK("https://example.invalid")'], 'test')
            path = Path(temp) / 'result.xlsx'
            write_report(report, path, 'test')
            workbook = openpyxl.load_workbook(path, data_only=False)
            self.assertEqual(workbook['KyQuanTri']['B2'].data_type, 's')
            workbook.close()


SOURCE = ROOT / 'MonthlyReportingData'
REFERENCE = Path(os.environ.get('REPORT_REFERENCE_FILE', 'D:/Dashboard/DuLieuBaoCao.xlsx'))


@unittest.skipUnless(SOURCE.exists() and REFERENCE.exists(),
                     'Kiểm tra đối chiếu cần bộ nguồn gốc và file mẫu (REPORT_REFERENCE_FILE).')
class IntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.output = Path(cls.temp.name) / 'result.xlsx'
        with contextlib.redirect_stdout(io.StringIO()):
            cls.result = run_report(Settings(SOURCE, cls.output, year=2026))
        cls.actual = openpyxl.load_workbook(cls.output, data_only=True)
        cls.reference = openpyxl.load_workbook(REFERENCE, data_only=True)
        cls.audit = json.loads(cls.result.audit_file.read_text(encoding='utf-8'))

    @classmethod
    def tearDownClass(cls):
        cls.actual.close()
        cls.reference.close()
        cls.temp.cleanup()

    def test_all_sheets_match_reference(self):
        """Khớp mẫu theo tiêu đề, thứ tự dòng và tên bảng Excel mà Power BI đọc.

        TaiChinhNhaMay: mẫu làm tròn nửa đồng về số chẵn (26.567.349.356,5 → …356), chương trình
        làm tròn lên như mọi sheet khác (…357), nên chỉ cột GiaTri được lệch tối đa 1 đồng."""
        def normalized(value):
            return round(value, 6) if isinstance(value, float) else value
        self.assertEqual(self.actual.sheetnames, self.reference.sheetnames)
        for name, headers in load_config('schemas.json').items():
            with self.subTest(sheet=name):
                def records(workbook):
                    return [tuple(normalized(v) for v in row) for row in
                            workbook[name].iter_rows(max_col=len(headers), values_only=True)]
                actual, expected = records(self.actual), records(self.reference)
                if name == 'TaiChinhNhaMay':
                    self.assertEqual([r[:3] for r in actual], [r[:3] for r in expected])
                    self.assertTrue(all(abs(a[3] - e[3]) <= 1 for a, e in zip(actual[1:], expected[1:])))
                else:
                    self.assertEqual(actual, expected)
                self.assertEqual(dict(self.actual[name].tables.items()), dict(self.reference[name].tables.items()))
        def product_list(workbook):
            return [row[0] for row in workbook['GhiChu'].iter_rows(values_only=True)]
        self.assertEqual(product_list(self.actual), product_list(self.reference))

    def test_checks_counts_and_provenance(self):
        self.assertEqual(self.result.business_rows, 2804)
        self.assertTrue(all(item['passed'] for item in self.audit['checks']))
        self.assertEqual(len(self.audit['sources']), 16)
        self.assertEqual(self.audit['status'], 'completed')
        for name, count in self.audit['row_counts'].items():
            self.assertEqual(len(self.audit['lineage'][name]), count)

    def test_plant_information_contains_monthly_and_cumulative_values(self):
        sheet = self.actual['ThongTinNhaMay']
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 416)
        self.assertEqual({row[1] for row in rows},
                         {f'T{month}' for month in range(1, 9)} |
                         {f'LK{month}T' for month in range(1, 9)})
        values = {(row[0], row[1], row[2]): row[3] for row in rows}
        self.assertEqual(values['SPH-1', 'T1', 'Quy mô sản xuất thiết kế'], 21600)
        self.assertEqual(values['SPH-2', 'T1', 'Quy mô sản xuất thiết kế'], 22800)
        self.assertEqual(values['SPH-1', 'LK1T', 'Sản lượng sản xuất'], 499.5752)
        self.assertEqual(values['SPH-2', 'LK1T', 'Sản lượng sản xuất'], 493.0668)
        self.assertEqual(sum(metric == 'Sản lượng điện tiêu thụ'
                     for _, _, metric in values), 32)
        # Tiêu hao nguyên liệu: 2 chỉ tiêu × 2 nhà máy × 16 kỳ, xếp cuối sheet như mẫu.
        self.assertEqual(round(values['SPH-1', 'T1', 'Bông tiêu hao'], 7), 398.3256626)
        self.assertEqual(round(values['SPH-2', 'T1', 'Xơ tiêu hao'], 7), 205.1028944)
        self.assertEqual({row[2] for row in rows[-64:]}, {'Bông tiêu hao', 'Xơ tiêu hao'})

    def test_plant_finance(self):
        rows = list(self.actual['TaiChinhNhaMay'].iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 128)
        values = {(row[0], row[1], row[2]): row[3] for row in rows}
        self.assertEqual(values['SPH-1', 'T1', 'Doanh thu sản xuất'], 31968260433)
        self.assertEqual(values['SPH-2', 'T1', 'Lợi nhuận trước thuế'], 1083954371)
        # Số dư TSCĐ chỉ có ở cột lũy kế nguồn: Tn và LKnT cùng giá trị cuối kỳ.
        self.assertEqual(values['SPH-2', 'T1', 'Nguyên giá TSCĐ'], 267569963525)
        self.assertEqual(values['SPH-1', 'T2', 'Giá trị còn lại TSCĐ'],
                         values['SPH-1', 'LK2T', 'Giá trị còn lại TSCĐ'])

    def test_sales_summary_units_table_and_adjusted_cumulative(self):
        sheet = self.actual['TongHopKQBanHang']
        self.assertEqual(sheet.max_row, 145)
        self.assertIn('TongHopKQBanHangTable', sheet.tables)
        self.assertEqual(sheet['D2'].number_format, '#,##0.######')
        self.assertEqual(sheet['D6'].number_format, '#,##0')
        values = {(r[1], r[2]): r[3] for r in sheet.iter_rows(min_row=2, values_only=True)}
        # LK4T gồm điều chỉnh nguồn, không bằng tổng doanh thu tháng 1–4.
        self.assertEqual(values['LK4T', 'Doanh thu nội địa'], 32118491933)
        self.assertNotEqual(values['LK4T', 'Doanh thu nội địa'],
                            sum(values[f'T{m}', 'Doanh thu nội địa'] for m in range(1, 5)))
        self.assertEqual(values['T2', 'Lợi nhuận trước thuế'], -2170725673)
        self.assertEqual(values['T8', 'Số lượng bán nội địa'], 11.5)
        self.assertTrue(any(w['code'] == 'SALES_SUMMARY_CUMULATIVE_ADJUSTMENT'
                            for w in self.audit['warnings']))

    def test_sales_summary_rejects_missing_amount(self):
        month = discover(SOURCE, 2026)[0]
        report = Report(load_config('schemas.json'))
        with Source(month.sales, 2026, 1, report) as source:
            source.sheet('1. KQ SXKD')['V6'] = None
            source.formulas['1. KQ SXKD']['V6'] = None
            with self.assertRaisesRegex(DataError, 'Cần số hợp lệ'):
                sales_summary.extract(source, report)

    def test_sales_summary_rejects_forecast_and_wrong_unit(self):
        month = discover(SOURCE, 2026)[1]  # T2 có cả cột thực hiện và ước quý.
        for cell, replacement in [('V4', 'ƯỚC THỰC HIỆN QUÝ 1/2026'),
                                  ('V5', 'Tháng báo cáo (đồng)')]:
            with self.subTest(cell=cell):
                report = Report(load_config('schemas.json'))
                with Source(month.sales, 2026, 2, report) as source:
                    source.sheet('1. KQ SXKD')[cell] = replacement
                    with self.assertRaisesRegex(DataError, 'đúng một cột thực hiện'):
                        sales_summary.extract(source, report)

    def test_prior_december_data_outside_catalog_and_known_exceptions(self):
        # T12/LK12T năm trước có trong dữ liệu chi phí nhưng không vào danh mục kỳ của Power BI.
        periods = {r[0] for r in self.actual['KyQuanTri'].iter_rows(min_row=2, values_only=True)}
        self.assertNotIn('T12', periods)
        self.assertNotIn('LK12T', periods)
        cost_periods = {r[4] for r in self.actual['KetQuaSanXuat'].iter_rows(min_row=2, values_only=True)}
        self.assertTrue({'T12', 'LK12T'}.issubset(cost_periods))
        codes = {w['code'] for w in self.audit['warnings']}
        self.assertTrue({'REFERENCE_PLAN_OVERRIDE', 'REFERENCE_CONTRACT_DATE',
                         'LABOR_CUMULATIVE_OUTLIER', 'KNOWN_INVENTORY_REF_ZERO'}.issubset(codes))

    def test_output_formulas_and_cached_results(self):
        workbook = openpyxl.load_workbook(self.output, data_only=False)
        self.assertEqual(workbook['SanLuongSanXuat']['B2'].data_type, 'f')
        self.assertIn('SanLuongSanXuat[[#This Row],[MatHang]]', workbook['SanLuongSanXuat']['B2'].value)
        self.assertEqual(workbook['HDBong']['E2'].value,
                         '=MuaBongTable[[#This Row],[SoLuongTan]]*1000*MuaBongTable[[#This Row],[DonGiaUSDkg]]')
        self.assertIn('SORT(', workbook['GhiChu']['A2'].value.text)
        self.assertEqual(self.actual['SanLuongSanXuat']['B2'].value, 'CM/Modal')
        self.assertEqual(self.actual['HDBong']['E2'].value, 860000)
        workbook.close()


class ReceivableAssembleTests(unittest.TestCase):
    """Ghép nước qua các tháng bằng dữ liệu giả lập, không cần file nguồn."""
    NAMES = receivables.VALUE_NAMES

    def month(self, number, countries):
        values = lambda **kw: {name: kw.get(name, 0.0) for name in self.NAMES}
        segments = {receivables.DOMESTIC: (values(), 'nd')}
        for name, kw in countries.items():
            segments[name, 'Xuất khẩu'] = (values(**kw), name)
        segments[receivables.ONSITE] = (values(), 'xktc')
        total = {name: sum(v[0][name] for v in segments.values()) for name in self.NAMES}
        return dict(year=2026, month=number, segments=segments,
                    total=receivables._convert(total), path=f'T{number}.xlsx')

    def assemble(self, months):
        report = Report(load_config('schemas.json'))
        receivables.assemble(months, report)
        return report, {(r[0], r[1]): r[3:] for r in report.rows['CongNoPhaiThu']}

    def test_country_first_seen_later_is_filled_for_every_month(self):
        months = [self.month(1, {'Nhật Bản': {}, 'Ai Cập': {}}),
                  self.month(2, {'Nhật Bản': {}, 'Ai Cập': {}}),
                  self.month(3, {'Nhật Bản': dict(DoanhSoThangDong=1.5, DoanhSoLuyKeDong=1.5), 'Ai Cập': {}})]
        _, rows = self.assemble(months)
        self.assertEqual(len(rows), 9)  # 3 tháng × (Việt Nam, Nhật Bản, XKTC); Ai Cập luôn 0.
        self.assertEqual(rows['T1', 'Nhật Bản'], [0] * 9)
        self.assertEqual(rows['T3', 'Nhật Bản'][1:3], [1500000, 1500000])

    def test_balance_only_country_is_kept(self):
        _, rows = self.assemble([self.month(1, {'Malaysia': dict(SoDu=0.2, TrongHan=0.2)}),
                                 self.month(2, {'Malaysia': {}})])
        self.assertEqual(rows['T1', 'Malaysia'][3], 200000)
        self.assertEqual(rows['T2', 'Malaysia'], [0] * 9)

    def test_missing_row_with_prior_cumulative_is_error(self):
        months = [self.month(1, {'Macau': dict(DoanhSoThangDong=2, DoanhSoLuyKeDong=2)}),
                  self.month(2, {})]
        with self.assertRaisesRegex(DataError, 'thiếu dòng nước Macau'):
            self.assemble(months)

    def test_missing_row_without_prior_value_is_zero_with_warning(self):
        report, rows = self.assemble([self.month(1, {}), self.month(2, {'USA': dict(SoDu=1, TrongHan=1)})])
        self.assertEqual(rows['T1', 'USA'], [0] * 9)
        self.assertTrue(any(w['code'] == 'RECEIVABLE_COUNTRY_ROW_MISSING' for w in report.warnings))

    def test_country_alias(self):
        self.assertEqual(country_name('  Thổ  Nhỹ Kỳ '), 'Thổ Nhĩ Kỳ')
        self.assertEqual(country_name('Hàn Quốc'), 'Hàn Quốc')


@unittest.skipUnless(SOURCE.exists(), 'Cần bộ nguồn MonthlyReportingData.')
class ReceivableTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.report = Report(load_config('schemas.json'))
        months = []
        for item in discover(SOURCE, 2026):
            with Source(item.sales, 2026, item.month, cls.report) as source:
                months.append(receivables.extract(source, cls.report))
        receivables.assemble(months, cls.report)
        cls.rows = {(r[0], r[1], r[2]): r[3:] for r in cls.report.rows['CongNoPhaiThu']}

    def extract(self, month, change):
        item = discover(SOURCE, 2026)[month - 1]
        report = Report(load_config('schemas.json'))
        with Source(item.sales, 2026, month, report) as source:
            change(source)
            receivables.extract(source, report)

    def test_full_grid_for_every_month(self):
        countries = {(c, k) for _, c, k in self.rows}
        self.assertEqual(countries, {('Việt Nam', 'Nội địa'), ('Xuất khẩu tại chỗ', 'Xuất khẩu'),
                                     *((c, 'Xuất khẩu') for c in ('Hàn Quốc', 'Hồng Kông', 'Malaysia',
                                                                  'Nhật Bản', 'Philippines', 'Trung Quốc',
                                                                  'Singapore'))})
        self.assertEqual(len(self.rows), 8 * len(countries))
        self.assertTrue(all(c['passed'] for c in self.report.checks))

    def test_values_in_dong_and_intermittent_countries(self):
        self.assertEqual(self.rows['T4', 'Việt Nam', 'Nội địa'][:4], [91.344, 9793000000, 32118491933, 8487261082])
        self.assertEqual(self.rows['T4', 'Trung Quốc', 'Xuất khẩu'][:4], [398.75, 27953000000, 60155174919, 23829000000])
        # Malaysia: T2 chỉ có số dư mang sang, T3 không phát sinh, T4 có bán.
        self.assertEqual(self.rows['T2', 'Malaysia', 'Xuất khẩu'][3], 184631228)
        self.assertEqual(self.rows['T3', 'Malaysia', 'Xuất khẩu'], [0] * 9)
        self.assertEqual(self.rows['T4', 'Malaysia', 'Xuất khẩu'][:3], [23.66, 1543000000, 1543000000])
        # Nhật Bản bán lần đầu T3: T1, T2 vẫn có dòng 0; T8 giữ lũy kế.
        self.assertEqual(self.rows['T1', 'Nhật Bản', 'Xuất khẩu'], [0] * 9)
        self.assertEqual(self.rows['T8', 'Nhật Bản', 'Xuất khẩu'][2], 1282852968)
        month4 = [v for (ky, _, _), v in self.rows.items() if ky == 'T4']
        self.assertEqual(sum(v[3] for v in month4), 48337261082)
        self.assertEqual(sum(v[1] for v in month4), 64448000000)
        codes = {w['code'] for w in self.report.warnings}
        self.assertIn('COUNTRY_NAME_ALIAS', codes)
        self.assertIn('RECEIVABLE_COUNTRY_LIST_CHANGED', codes)

    def test_detail_must_match_group(self):
        def tamper(source):
            source.sheet('2. Phai thu')['I28'] = 5000
        with self.assertRaisesRegex(DataError, 'Xuất khẩu/TrongHan/tổng chi tiết'):
            self.extract(4, tamper)

    def test_number_without_sequence_is_rejected(self):
        def tamper(source):
            sheet = source.sheet('2. Phai thu')
            sheet['A30'] = None  # Indonesia mất STT nhưng có số dư.
            sheet['H30'] = 1
        with self.assertRaisesRegex(DataError, 'không có STT'):
            self.extract(4, tamper)

    def test_aging_must_match_balance(self):
        def tamper(source):
            source.sheet('2. Phai thu')['J9'] = 5
        with self.assertRaisesRegex(DataError, 'số dư = tổng tuổi nợ'):
            self.extract(4, tamper)

    def test_moved_header_is_rejected(self):
        def tamper(source):
            source.sheet('2. Phai thu')['I8'] = 'Chưa đến hạn'
        with self.assertRaisesRegex(DataError, 'Trong hạn'):
            self.extract(4, tamper)

if __name__ == '__main__':
    unittest.main()
