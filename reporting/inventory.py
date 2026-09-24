"""Tồn kho Bông/Xơ tổng hợp, chi tiết theo nguồn gốc và tồn kho thành phẩm."""
from openpyxl.utils import get_column_letter
from .common import key, text, rounded, dong
from .models import DataError

YARN_TYPES=['Sợi Cotton chải thô','Sợi Cotton chải kỹ','Sợi CVC chải thô',
            'Sợi CVC chải kỹ','Sợi TC chải thô','Sợi TC chải kỹ','Sợi PE','Sợi khác']
ORIGINS={'bema':'Benma 10','SENFA':'Shengfa','Paraguary':'Paraguay'}
STATES=['Tồn kho cuối tháng','Tồn kho trên 3 tháng','Giảm chất lượng']


def extract(source, report, overrides):
    _materials(source, report, overrides['inventory_zero_errors'])
    _finished_goods(source, report)


def _materials(source, report, exception):
    """Ghi chú triển khai:
- Phân biệt dòng tổng nguyên liệu (không có nguồn gốc ở E) và chi tiết.
- Phần tổng hợp: E là số nhập kho, phần chi tiết: E là tên nguồn gốc."""
    sheet = '2a. HTK NL'
    sh = source.sheet(sheet)
    summaries = {}
    details = []
    detail_header = source.find(sheet, 'Nguồn gốc', 5)
    for row in range(1, sh.max_row + 1):
        material = text(sh.cell(row, 4).value)
        if material not in ('Bông', 'Xơ'):
            continue
        origin = text(sh.cell(row, 5).value)
        if row > detail_header:
            if origin:
                details.append((row, material, ORIGINS.get(origin, origin)))
            elif any((source.number(sheet, row, c, blank_zero=True) for c in [6, 9, 10, 13, 14, 18])):
                raise DataError(f'Có tồn kho nhưng thiếu nguồn gốc: {source.loc(sheet, row, 5)}')
        else:
            if material in summaries:
                raise DataError(f'Lặp dòng tổng {material}: {source.path}')
            summaries[material] = row
    if set(summaries) != {'Bông', 'Xơ'} or not details:
        raise DataError(f'Không nhận diện bảng nguyên liệu: {source.path}')
    aggregates = {material: [[0.0, 0.0] for _ in STATES] for material in summaries}
    for row, material, origin in details:

        def num(col):
            value = source.value(sheet, row, col)
            address = f'{get_column_letter(col)}{row}'
            if isinstance(value, str) and value.startswith('#'):
                allowed = source.year == exception['year'] and source.month == exception['month'] and (address in exception['cells']) and (value == exception['expected_error']) and (source.number(sheet, row, 13, blank_zero=True) == 0)
                if not allowed:
                    raise DataError(f'Lỗi Excel: {source.loc(sheet, row, col)} = {value}')
                report.warn('KNOWN_INVENTORY_REF_ZERO', exception['reason'], source.loc(sheet, row, col))
                return 0.0
            return source.number(sheet, row, col, blank_zero=True)
        for index, (qty_col, usd_col, amount_col) in enumerate([(6, 8, 9), (10, 12, 13), (14, 16, 18)]):
            qty = num(qty_col)
            usd = num(usd_col)
            amount = num(amount_col)
            aggregates[material][index][0] += qty
            aggregates[material][index][1] += amount
            report.add('TonKhoNguyenLieuChiTiet', [f'T{source.month}', material, origin, STATES[index], rounded(qty, 6), rounded(usd, 6), dong(amount)], source.loc(sheet, row))
    in_process = any(('tồn trên chuyền' in key(c.value) for row in sh.iter_rows(max_row=min(summaries.values()) - 1) for c in row if c.column == 13))
    for material, row in summaries.items():
        final = (source.number(sheet, row, 7, blank_zero=True), source.number(sheet, row, 10, blank_zero=True))
        for index in [0, 1]:
            report.check(f'T{source.month}/{material}/tồn cuối/{index}', aggregates[material][0][index], final[index], 1e-06, source.loc(sheet, row))
        transit = (source.number(sheet, row, 11, blank_zero=True), source.number(sheet, row, 12, blank_zero=True))
        work = (source.number(sheet, row, 13, blank_zero=True), source.number(sheet, row, 14, blank_zero=True)) if in_process else (0, 0)
        if not in_process:
            report.warn('NO_WORK_IN_PROGRESS_COLUMNS', f'T{source.month}: chưa có cột tồn trên chuyền; dùng 0 theo mẫu.', source.loc(sheet, row))
        for state, pair in [(STATES[0], final), ('Đi đường và chốt mua', transit), (STATES[1], aggregates[material][1]), (STATES[2], aggregates[material][2]), ('Tồn trên chuyền', work)]:
            report.add('TonKhoNguyenLieu', [f'T{source.month}', material, state, rounded(pair[0], 6), dong(pair[1])], source.loc(sheet, row) + '; chi tiết theo nguồn gốc')


def _finished_goods(source, report):
    """Ghi chú triển khai:
- Cùng loại sợi xuất hiện ở bảng giá phía dưới: lấy khối đầu tiên."""
    sheet = '2b. HTK Soi'
    sh = source.sheet(sheet)
    starts = [r for r in range(1, sh.max_row + 1) if key(sh.cell(r, 4).value) == key(YARN_TYPES[0])]
    if not starts:
        raise DataError(f'Thiếu khối tồn kho sợi: {source.path}')
    start = starts[0]
    sums = [0.0] * 4
    for offset, yarn in enumerate(YARN_TYPES):
        row = start + offset
        source.expect(sheet, row, 4, yarn)
        # Xuất khác (cột 6) không xuất ra mẫu nhưng vẫn dùng để đối chiếu dòng tổng.
        vals = [source.number(sheet, row, c, blank_zero=True) for c in [5, 6, 8, 11]]
        report.add('TonKhoThanhPham', [f'T{source.month}', yarn, rounded(vals[0], 6), rounded(vals[2], 6), dong(vals[3])], source.loc(sheet, row))
        for i, value in enumerate(vals):
            sums[i] += value
    total = start + len(YARN_TYPES)
    if 'tổng' not in key(sh.cell(total, 4).value):
        raise DataError(f'Thiếu tổng tồn kho sợi: {source.loc(sheet, total)}')
    for i, col in enumerate([5, 6, 8, 11]):
        report.check(f'T{source.month}/Tồn kho sợi/cột {col}', sums[i], source.number(sheet, total, col, blank_zero=True), 1e-06, source.loc(sheet, total, col))
