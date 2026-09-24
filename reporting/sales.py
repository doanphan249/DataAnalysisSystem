"""Giá tiêu thụ sợi, phế liệu thanh lý, hợp đồng bông và tiêu thụ theo nước."""
from datetime import datetime, date
from openpyxl.utils.datetime import from_excel
from .common import country_name, key, text, rounded, dong
from .inventory import YARN_TYPES
from .models import DataError

WASTE={
 'Bông phế F1':'Bông phế từ công đoạn máy bông',
 'Bông phế F2':'Bông phế từ công đoạn chải thô',
 'Bông phế F1 + F2':'Bông phế công đoạn máy bông - chải thô',
 'Bông phế F3':'Bông phế từ công đoạn chải kỹ',
 'Bông phế khác':'Bông phế khác',
 'Sợi rối':'Sợi phế từ công đoạn máy ống',
 'Sợi phế các loại':'Sợi phế từ các công đoạn khác',
 'Xơ phế':'Xơ phế thu hồi tại các công đoạn',
}


def extract_production_book(source, report, overrides, apply_overrides):
    _yarn_prices(source, report)
    _waste(source, report)
    _contracts(source, report, overrides, apply_overrides)


def _yarn_prices(source, report):
    sheet = '2b. HTK Soi'
    sh = source.sheet(sheet)
    header = source.find(sheet, 'Mặt hàng', 5)
    end = source.find(sheet, 'Bông phế F1', 4)
    count = 0
    for row in range(header + 1, end):
        yarn = text(sh.cell(row, 4).value)
        if yarn not in YARN_TYPES:
            continue
        item = sh.cell(row, 5).value
        if item in (None, '', '-') or item == 0:
            continue
        if not isinstance(item, str):
            raise DataError(f'Sai tên mặt hàng: {source.loc(sheet, row, 5)}')
        # Mẫu chỉ dùng giá CIF (cột 9); giá nội địa/EXW/FOB không xuất.
        price = source.number(sheet, row, 9, blank_zero=True)
        if source.value(sheet, row, 9) in (None, '', '-'):
            report.warn('EMPTY_SELLING_PRICE_ZERO', 'Giá chưa có được biểu diễn bằng 0 theo mẫu; không phải giá giao dịch bằng 0.', source.loc(sheet, row, 9))
        quality = []
        for col in [10, 11]:
            value = source.value(sheet, row, col)
            quality.append(None if value in (None, '', '-') else source.number(sheet, row, col))
        report.add('TieuThuSoi', [f'T{source.month}', yarn, text(item), price, *quality], source.loc(sheet, row))
        count += 1
    if not count:
        report.warn('NO_YARN_PRICE_ROWS', f'T{source.month} không có mặt hàng báo giá.', str(source.path))


def _waste(source, report):
    sheet = '2b. HTK Soi'
    start = source.find(sheet, next(iter(WASTE)), 4)
    sums = [0.0, 0.0]
    for offset, (original, target) in enumerate(WASTE.items()):
        row = start + offset
        source.expect(sheet, row, 4, original)
        qty = source.number(sheet, row, 5, blank_zero=True)
        amount = source.number(sheet, row, 7, blank_zero=True)
        report.add('PheLieuThanhLy', [f'T{source.month}', target, qty, amount], source.loc(sheet, row))
        sums[0] += qty
        sums[1] += amount
    total = start + len(WASTE)
    if 'tổng' not in key(source.sheet(sheet).cell(total, 4).value):
        raise DataError(f'Thiếu tổng phế liệu: {source.loc(sheet, total)}')
    for index, col, tol in [(0, 5, 1e-06), (1, 7, 0.5)]:
        report.check(f'T{source.month}/Phế liệu/cột {col}', sums[index], source.number(sheet, total, col), tol, source.loc(sheet, total, col))


def _contracts(source, report, overrides, apply_overrides):
    """Ghi chú triển khai:
- Hàng hợp đồng phải có STT số để không nhập ghi chú cuối trang."""
    sheet = '4. HD bong'
    sh = source.sheet(sheet)
    headings = [source.find(sheet, name, 3) for name in ['Ghi giảm kỳ trước', 'Phát sinh trong kỳ']]
    for index, name in enumerate(['Ghi giảm kỳ trước', 'Phát sinh trong kỳ']):
        start = headings[index]
        end = headings[index + 1] if index == 0 else sh.max_row + 1
        total_qty = total_amount = 0.0
        for row in range(start + 1, end):
            seq = sh.cell(row, 1).value
            if not isinstance(seq, (int, float)):
                continue
            qty = source.number(sheet, row, 5, blank_zero=True)
            amount = source.number(sheet, row, 7, blank_zero=True)
            if qty == 0 and amount == 0:
                continue
            price = source.number(sheet, row, 6)
            origin = text(sh.cell(row, 4).value)
            if not origin:
                raise DataError(f'Thiếu nguồn gốc hợp đồng: {source.loc(sheet, row, 4)}')
            raw_date = source.value(sheet, row, 3)
            contract_date = raw_date if isinstance(raw_date, (date, datetime)) else None
            if isinstance(raw_date, (int, float)):
                contract_date = from_excel(raw_date, source.values.epoch)
            if contract_date is None and apply_overrides:
                for override in overrides['contract_dates']:
                    identity = (source.year, source.month, origin, qty, price, raw_date)
                    expected = (override['year'], override['month'], override['origin'], override['quantity'], override['price'], override['expected_source'])
                    if identity == expected:
                        contract_date = datetime.fromisoformat(override['date'])
                        report.warn('REFERENCE_CONTRACT_DATE', override['reason'], source.loc(sheet, row, 3))
            if contract_date is None:
                report.warn('MISSING_CONTRACT_DATE', f'Giữ trống ngày không hợp lệ {raw_date!r}.', source.loc(sheet, row, 3))
            calculated = qty * 1000 * price
            report.check(f'T{source.month}/HĐ bông/hàng {row}', calculated, amount, 0.01, source.loc(sheet, row, 7))
            if name == 'Phát sinh trong kỳ':
                report.add('HDBong', [f'T{source.month}', contract_date, qty, price, calculated], source.loc(sheet, row))
            else:
                # MuaBongTable không có cột loại giao dịch: ghi giảm sẽ bị cộng như hợp đồng mua mới.
                report.warn('CONTRACT_REDUCTION_NOT_EXPORTED', f'{name}: {qty} tấn không xuất vào HDBong; cần xử lý riêng trên báo cáo.', source.loc(sheet, row))
            total_qty += qty
            total_amount += amount
        report.check(f'T{source.month}/HĐ bông/{name}/lượng', total_qty, source.number(sheet, start, 5, blank_zero=True), 1e-06, source.loc(sheet, start, 5))
        report.check(f'T{source.month}/HĐ bông/{name}/tiền', total_amount, source.number(sheet, start, 7, blank_zero=True), 0.01, source.loc(sheet, start, 7))


def read_country_sales(source, report) -> dict:
    """Trả dữ liệu tháng. Pipeline ghép sau khi biết danh sách nước của cả năm."""
    sheet = '2. Phai thu'
    sh = source.sheet(sheet)
    source.validate_date(sheet, 'D4')
    export = source.find(sheet, 'Thị trường xuất khẩu', 4)
    domestic = source.find(sheet, 'Thị trường nội địa', 4)
    grand = source.find(sheet, 'TỔNG CỘNG', 4)
    xktc = source.find(sheet, 'XKTC', 4)

    def metrics(row):
        return tuple((source.number(sheet, row, c, blank_zero=True) for c in [5, 6, 7]))
    countries = {}
    for row in range(export + 1, sh.max_row + 1):
        name = country_name(sh.cell(row, 4).value)
        if 'phần i -' in key(name):
            break
        seq = sh.cell(row, 1).value
        if name and name != 'XKTC' and isinstance(seq, (int, float)):
            if name in countries:
                raise DataError(f'Quốc gia trùng {name}: {source.path}')
            countries[name] = metrics(row)
    exp = metrics(export)
    dom = metrics(domestic)
    onsite = metrics(xktc)
    total = metrics(grand)
    for i in range(3):
        report.check(f'T{source.month}/Xuất khẩu/chỉ số {i}', sum((v[i] for v in countries.values())) + onsite[i], exp[i], 1e-06, source.loc(sheet, export))
        report.check(f'T{source.month}/Tổng tiêu thụ/chỉ số {i}', exp[i] + dom[i], total[i], 1e-06, source.loc(sheet, grand))
    return dict(year=source.year, month=source.month, countries=countries, domestic=dom, onsite=onsite, source=source.loc(sheet, 'D:G'))


def assemble_country_sales(months, report):
    # Nước có số khác 0 ở bất kỳ tháng nào, theo thứ tự dòng trên form như CongNoPhaiThu.
    sold = {c for month in months for c, values in month['countries'].items() if any(values)}
    active = []
    for month in months:
        active += [c for c in month['countries'] if c in sold and c not in active]
    segments = [(c, 'Xuất khẩu') for c in active] + [('Việt Nam', 'Nội địa'), ('Xuất khẩu tại chỗ', 'Xuất khẩu')]
    quantities = {s: 0.0 for s in segments}
    previous = {s: 0 for s in segments}
    for month in months:
        for country, channel in segments:
            segment = (country, channel)
            if country == 'Việt Nam' and channel == 'Nội địa':
                values = month['domestic']
            elif country == 'Xuất khẩu tại chỗ':
                values = month['onsite']
            else:
                if country not in month['countries'] and (quantities[segment] or previous[segment]):
                    raise DataError(f"T{month['month']} thiếu dòng nước {country} đã phát sinh trước đó; không thể giả định lũy kế bằng 0.")
                values = month['countries'].get(country, (0, 0, 0))
            qty, amount, cumulative = values
            quantities[segment] = rounded(quantities[segment] + qty, 6)
            difference = dong(cumulative) - previous[segment] - dong(amount)
            if abs(difference) > 1:
                report.warn('SALES_CUMULATIVE_ADJUSTMENT', f"T{month['month']} {country}/{channel}: lũy kế nguồn khác lũy kế trước + tháng {difference:,.0f} đồng. Giữ lũy kế nguồn.", month['source'])
            for ky, value in [(f"T{month['month']}", dong(amount)), (f"LK{month['month']}T", dong(cumulative))]:
                report.add('PhanTichTieuThu', [ky, country, channel, value], month['source'])
            previous[segment] = dong(cumulative)
