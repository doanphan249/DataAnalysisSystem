"""Tìm file, kiểm tra kỳ, đọc số đã tính và phát hiện công thức chưa có cache."""
from pathlib import Path
from datetime import datetime, date
import calendar
import math
import re
import warnings
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel
from .common import key, text
from .models import DataError, MonthFiles, Report


def discover(folder: Path, year: int | None) -> list[MonthFiles]:
    if not folder.is_dir():
        raise DataError(f'Không có thư mục nguồn: {folder}')
    directories = []
    for path in sorted(folder.iterdir()):
        match = re.fullmatch('BCQT_(\\d{2})(\\d{4})', path.name, re.IGNORECASE)
        if path.is_dir() and match:
            month, candidate_year = map(int, match.groups())
            if not 1 <= month <= 12:
                raise DataError(f'Tên thư mục có tháng sai: {path}')
            directories.append((candidate_year, month, path))
    years = {y for y, _, _ in directories}
    if year is None:
        if len(years) != 1:
            raise DataError(f'Nguồn có các năm {sorted(years)}. Đặt REPORT_YEAR trong main.py.')
        year = next(iter(years))
    selected = sorted(((m, p) for y, m, p in directories if y == year))
    if not selected:
        raise DataError(f'Không có dữ liệu năm {year}')
    months = [m for m, _ in selected]
    if months != list(range(1, max(months) + 1)):
        raise DataError(f'Cần đủ tháng từ T1 để tính sản lượng tiêu thụ lũy kế. Đang có: {months}')
    result = []
    for month, directory in selected:
        files = [p for p in directory.glob('*.xlsx') if not p.name.startswith('~$')]

        def select(token):
            found = [p for p in files if token in key(p.stem)]
            if len(found) != 1:
                raise DataError(f"Cần đúng 1 file '{token}' trong {directory}; thấy: {[p.name for p in found]}")
            return found[0]
        result.append(MonthFiles(year, month, select('qtsx'), select('kq tieu thu')))
    return result


class Source:
    """Mở hai chế độ: giá trị Excel đã lưu và công thức gốc. Không sửa nguồn."""

    def __init__(self, path: Path, year: int, month: int, report: Report):
        self.path, self.year, self.month, self.report = (path, year, month, report)
        with warnings.catch_warnings():
            warnings.filterwarnings('ignore', message='Data Validation extension is not supported.*')
            self.values = openpyxl.load_workbook(path, data_only=True)
            self.formulas = openpyxl.load_workbook(path, data_only=False)

    def __enter__(self):
        return self

    def __exit__(self, *_):
        self.values.close()
        self.formulas.close()

    def sheet(self, name: str):
        if name not in self.values.sheetnames:
            raise DataError(f"Thiếu sheet '{name}': {self.path}")
        return self.values[name]

    def loc(self, sheet: str, row: int | str, col: int | None=None) -> str:
        cell = f'{get_column_letter(col)}{row}' if col else str(row)
        return f'{self.path} :: {sheet}!{cell}'

    def value(self, sheet: str, row: int, col: int):
        cell = self.sheet(sheet).cell(row, col)
        if cell.value is None and self.formulas[sheet].cell(row, col).data_type == 'f':
            raise DataError(f'Công thức chưa có giá trị đã lưu: {self.loc(sheet, row, col)}. Mở nguồn bằng Excel, tính lại và lưu.')
        return cell.value

    def number(self, sheet: str, row: int, col: int, *, blank_zero: bool=False) -> float:
        value = self.value(sheet, row, col)
        if (value is None or (isinstance(value, str) and text(value) in ('', '-'))) and blank_zero:
            return 0.0
        if isinstance(value, bool) or not isinstance(value, (int, float)) or (not math.isfinite(value)):
            raise DataError(f'Cần số hợp lệ, nhận {value!r}: {self.loc(sheet, row, col)}')
        return float(value)

    def find(self, sheet: str, label: str, col: int, start: int=1, end: int | None=None) -> int:
        sh = self.sheet(sheet)
        hits = [r for r in range(start, (end or sh.max_row) + 1) if key(sh.cell(r, col).value) == key(label)]
        if len(hits) != 1:
            raise DataError(f'Nhãn {label!r} cần duy nhất, tìm thấy {hits}: {self.path} :: {sheet}')
        return hits[0]

    def expect(self, sheet: str, row: int, col: int, label: str) -> None:
        actual = text(self.sheet(sheet).cell(row, col).value)
        if key(actual) != key(label):
            raise DataError(f'Bố cục thay đổi: cần {label!r}, nhận {actual!r}: {self.loc(sheet, row, col)}')

    def validate_date(self, sheet: str, cell: str) -> None:
        c = self.sheet(sheet)[cell]
        value = self.value(sheet, c.row, c.column)
        if isinstance(value, (float, int)):
            value = from_excel(value, self.values.epoch)
        expected = date(self.year, self.month, calendar.monthrange(self.year, self.month)[1])
        actual = value.date() if isinstance(value, datetime) else value
        if actual != expected:
            raise DataError(f'Ngày báo cáo {actual!r} khác {expected}: {self.loc(sheet, cell)}')

    def validate_company(self, sheet: str) -> None:
        title = key(self.sheet(sheet)['A1'].value)
        if title.startswith('chọn đơn vị báo cáo') and 'phú hưng' in key(self.path.name):
            self.report.warn('COMPANY_PLACEHOLDER', 'Tên công ty trong sheet còn là ký hiệu mẫu; xác định SPH từ tên file Phú Hưng.', self.loc(sheet, 'A1'))
            return
        if 'phú hưng' not in title:
            raise DataError(f'Chương trình được cấu hình cho SPH/Phú Hưng; tên đơn vị khác: {self.path} :: {sheet}!A1')

    def ensure_unit(self, sheet: str, row: int, col: int, expected: str) -> None:
        self.expect(sheet, row, col, expected)
